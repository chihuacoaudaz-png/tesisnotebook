# -*- coding: utf-8 -*-
import os, sys, docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import win32com.client

def build_report():
    print('[*] Generando informe Cuajone 2025 (RR 1439-2023) - Texto Negro & Figuras Limpias...')
    doc = docx.Document()
    
    # Margenes Oficiales UNI
    for s in doc.sections:
        s.top_margin = Cm(2.54)
        s.bottom_margin = Cm(2.50)
        s.left_margin = Cm(3.00)
        s.right_margin = Cm(2.50)
        s.different_first_page_header_footer = True
        
        hp = s.header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hr = hp.add_run('UNIVERSIDAD NACIONAL DE INGENIERÍA | FIGMM')
        hr.font.name = 'Arial'
        hr.font.size = Pt(8.5)
        hr.font.color.rgb = RGBColor(0, 0, 0)
        
        fp = s.footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        fr = fp.add_run('Trabajo de Investigación Técnica - Gerencia Mina Cuajone 2025')
        fr.font.name = 'Arial'
        fr.font.size = Pt(8.0)
        fr.font.color.rgb = RGBColor(0, 0, 0)

    # Estilos Base (100% Negro)
    norm = doc.styles['Normal']
    norm.font.name = 'Arial'
    norm.font.size = Pt(11)
    norm.font.color.rgb = RGBColor(0, 0, 0)

    def add_h1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(14)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(13)
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_h2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(text)
        r.bold = True
        r.font.size = Pt(11.5)
        r.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_body(text, bold_prefix='', italic=False):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.color.rgb = RGBColor(0, 0, 0)
        rt = p.add_run(text)
        rt.font.color.rgb = RGBColor(0, 0, 0)
        rt.italic = italic
        return p

    def add_bullet(text, bold_prefix=''):
        p = doc.add_paragraph(style='List Bullet')
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold_prefix:
            rb = p.add_run(bold_prefix)
            rb.bold = True
            rb.font.color.rgb = RGBColor(0, 0, 0)
        rt = p.add_run(text)
        rt.font.color.rgb = RGBColor(0, 0, 0)
        return p

    def add_formula(eq_text, desgloses=[]):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(eq_text)
        r.font.name = 'Consolas'
        r.font.size = Pt(10.0)
        r.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)
        if desgloses:
            pd = doc.add_paragraph()
            pd.paragraph_format.space_after = Pt(2)
            rd = pd.add_run('Donde:')
            rd.italic = True
            rd.font.size = Pt(9.5)
            rd.font.color.rgb = RGBColor(0, 0, 0)
            for d in desgloses:
                pi = doc.add_paragraph(style='List Bullet')
                pi.paragraph_format.space_after = Pt(2)
                pi.paragraph_format.line_spacing = 1.05
                ri = pi.add_run(d)
                ri.font.size = Pt(9.5)
                ri.font.color.rgb = RGBColor(0, 0, 0)

    def add_fig(img_path, caption_text, width_cm=10.5):
        if os.path.exists(img_path):
            pimg = doc.add_paragraph()
            pimg.paragraph_format.space_before = Pt(6)
            pimg.paragraph_format.space_after = Pt(3)
            pimg.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pimg.add_run().add_picture(img_path, width=Cm(width_cm))
            
            pcap = doc.add_paragraph()
            pcap.paragraph_format.space_after = Pt(2)
            pcap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rcap = pcap.add_run(caption_text)
            rcap.bold = True
            rcap.font.size = Pt(9.0)
            rcap.font.color.rgb = RGBColor(0, 0, 0)
            
            psrc = doc.add_paragraph()
            psrc.paragraph_format.space_after = Pt(8)
            psrc.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rsrc = psrc.add_run('Fuente: Registro de Operaciones Mina Cuajone (2025).')
            rsrc.italic = True
            rsrc.font.size = Pt(8.0)
            rsrc.font.color.rgb = RGBColor(0, 0, 0)
